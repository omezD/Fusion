from rest_framework.views import APIView
from rest_framework.response import Response
from django.shortcuts import get_object_or_404, redirect, render
from rest_framework import status,permissions
from rest_framework.decorators import api_view,permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.authentication import TokenAuthentication
from django.contrib.auth.models import User
from django.http import JsonResponse,HttpResponse
from django.core.files.storage import FileSystemStorage
from django.views.decorators.csrf import csrf_exempt
from django.utils.decorators import method_decorator
from ..models import *
from applications.academic_information.models import Student
from applications.globals.models import ExtraInfo, DepartmentInfo
from .serializers import PlacementScheduleSerializer, NotifyStudentSerializer
from applications.academic_information.api.serializers import StudentSerializers
from datetime import datetime
import io
from reportlab.pdfgen import canvas
from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,ListFlowable,ListItem
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.graphics.shapes import Line,Drawing
from rest_framework.test import APIRequestFactory


@method_decorator(csrf_exempt, name='dispatch')
@permission_classes([IsAuthenticated])
class PlacementScheduleView(APIView):

    def get(self, request, id=None): 
        if id:
            try:
                notify_schedule = NotifyStudent.objects.get(id=id)
                placement_schedule = PlacementSchedule.objects.get(notify_id=notify_schedule)
                combined_entry = {**NotifyStudentSerializer(notify_schedule).data, **PlacementScheduleSerializer(placement_schedule).data}
                return Response(combined_entry, status=status.HTTP_200_OK)
            except NotifyStudent.DoesNotExist:
                return Response({"error": "NotifyStudent not found"}, status=status.HTTP_404_NOT_FOUND)
            except PlacementSchedule.DoesNotExist:
                return Response({"error": "PlacementSchedule not found"}, status=status.HTTP_404_NOT_FOUND)
        else:
            combined_data = []
            notify_students = NotifyStudent.objects.all()
            for notify in notify_students:
                placements = PlacementSchedule.objects.filter(notify_id=notify.id)
                placement_serializer = PlacementScheduleSerializer(placements, many=True)
                notify_data = NotifyStudentSerializer(notify).data

                for placement in placement_serializer.data:
                    counting = StudentApplication.objects.filter(schedule_id_id=placement['id'],unique_id_id=request.user.username).count()
                    role_st = Role.objects.get(id=placement['role'])
                    check = True
                    if counting==0:
                        check=False
                    combined_entry = {**notify_data, **placement ,'check':check ,'role_st':role_st.role}
                    combined_data.append(combined_entry)
            
            return Response(combined_data, status=status.HTTP_200_OK)
        
    def post(self, request):
        placement_type = request.data.get("placement_type")
        company_name = request.data.get("company_name")
        ctc = request.data.get("ctc")
        description = request.data.get("description")
        schedule_at = request.data.get("schedule_at")
        date = request.data.get("placement_date")
        location = request.data.get("location")
        role = request.data.get("role")
        resume = request.FILES.get("resume")

        try:
            role_create, _ = Role.objects.get_or_create(role=role)
            notify = NotifyStudent.objects.create(
                placement_type=placement_type,
                company_name=company_name,
                description=description,
                ctc=ctc,
                timestamp=schedule_at,
            )

            schedule = PlacementSchedule.objects.create(
                notify_id=notify,
                title=company_name,
                description=description,
                placement_date=date,
                attached_file=resume,
                role=role_create,
                location=location,
                time=schedule_at,
            )

            return JsonResponse({"message": "Successfully Added Schedule"}, status=201)

        except Exception as e:
            return JsonResponse({"error": str(e)}, status=400)
    
    def delete(self, request, id):
        try:
            placement_schedule = PlacementSchedule.objects.get(id=id)
            notify_schedule = NotifyStudent.objects.get(id=placement_schedule.notify_id_id)
            notify_schedule.delete()
            placement_schedule.delete()

            return JsonResponse({"message": "Successfully Deleted"}, status=200)

        except Exception as e:
            return JsonResponse({"error": str(e)}, status=400)
        
    def put(self, request, id):
        try:
            placement_schedule = PlacementSchedule.objects.get(id=id)
            notify_schedule = NotifyStudent.objects.get(id=placement_schedule.notify_id_id)


            placement_type = request.data.get("placement_type", notify_schedule.placement_type)
            company_name = request.data.get("company_name", notify_schedule.company_name)
            ctc = request.data.get("ctc", notify_schedule.ctc)
            description = request.data.get("description", notify_schedule.description)
            schedule_at = request.data.get("schedule_at", notify_schedule.timestamp)
            date = request.data.get("placement_date", placement_schedule.placement_date)
            location = request.data.get("location", placement_schedule.location)
            role = request.data.get("role", placement_schedule.role)
            resume = request.FILES.get("resume", placement_schedule.attached_file)

            notify_schedule.placement_type = placement_type
            notify_schedule.company_name = company_name
            notify_schedule.ctc = ctc
            notify_schedule.description = description
            notify_schedule.timestamp = schedule_at
            notify_schedule.save()

            placement_schedule.title = company_name
            placement_schedule.description = description
            placement_schedule.placement_date = date
            placement_schedule.location = location
            placement_schedule.attached_file = resume
            placement_schedule.time = schedule_at
            role_create, _ = Role.objects.get_or_create(role=role)
            placement_schedule.role = role_create
            placement_schedule.save()

            return JsonResponse({"message": "Successfully Updated"}, status=200)
        
        except Exception as e:
            return JsonResponse({"error": str(e)}, status=400)
            
    

@permission_classes([IsAuthenticated]) 
class BatchStatisticsView(APIView):

    def get(self, request):
        combined_data = []
        student_records = StudentRecord.objects.all()

        if not student_records.exists():
            return Response({"error": "No student records found"}, status=status.HTTP_404_NOT_FOUND)

        for student in student_records:
            try:

                cur_student = Student.objects.get(id_id=student.unique_id_id)
                cur_placement = PlacementRecord.objects.get(id=student.record_id_id)
                user = User.objects.get(username=student.unique_id_id)

                combined_entry = {
                    "id": cur_placement.id ,
                    "branch": cur_student.specialization, 
                    "batch" : cur_placement.year, 

                    "placement_name": cur_placement.name,  
                    "ctc": cur_placement.ctc, 
                    "year": cur_placement.year, 
                    "first_name": user.first_name,
                }

                combined_data.append(combined_entry)

            except Student.DoesNotExist:
                return Response({"error": f"Student with id {student.unique_id} not found"}, status=status.HTTP_404_NOT_FOUND)
            except PlacementRecord.DoesNotExist:
                return Response({"error": f"Placement record with id {student.record_id} not found"}, status=status.HTTP_404_NOT_FOUND)
            except User.DoesNotExist:
                return Response({"error": f"User with id {student.unique_id} not found"}, status=status.HTTP_404_NOT_FOUND)
            except Exception as e:
                return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        if not combined_data:
            return Response({"message": "No combined data found"}, status=status.HTTP_204_NO_CONTENT)

        return Response(combined_data, status=status.HTTP_200_OK)



    def post(self,request):
        placement_type=request.POST.get("placement_type")
        company_name=request.POST.get("company_name")
        roll_no = request.POST.get("roll_no")
        ctc=request.POST.get("ctc")
        year=request.POST.get("year")
        test_type=request.POST.get("test_type")
        test_score=request.POST.get("test_score")

        try:
            p2 = PlacementRecord.objects.create(
                placement_type = placement_type,
                name = company_name,
                ctc = ctc,
                year = year,
                test_score = test_score,
                test_type = test_type,
            )
            p1 = StudentRecord.objects.create(
                record_id = p2,
                unique_id_id = roll_no,
            )
            return JsonResponse({"message": "Successfully Added"}, status=201)
    
        except Exception as e:
            return JsonResponse({"error": str(e)}, status=400)
    
    def put(self, request, record_id):
        try:
            placement_record = PlacementRecord.objects.get(id=record_id)
            placement_record.placement_type = request.data.get("placement_type", placement_record.placement_type)
            placement_record.name = request.data.get("company_name", placement_record.name)
            placement_record.ctc = request.data.get("ctc", placement_record.ctc)
            placement_record.year = request.data.get("year", placement_record.year)
            placement_record.test_score = request.data.get("test_score", placement_record.test_score)
            placement_record.test_type = request.data.get("test_type", placement_record.test_type)
            placement_record.save()

            return JsonResponse({"message": "Successfully Updated"}, status=200)

        except Exception as e:
            return JsonResponse({"error": str(e)}, status=400)
    

    def delete(self, request, id):
        try:
            placement_record = PlacementRecord.objects.get(id=id)
            student_record = StudentRecord.objects.get(record_id_id=id)
            student_record.delete()
            placement_record.delete()

            return JsonResponse({"message": "Successfully Deleted"}, status=200)

        except Exception as e:
            return JsonResponse({"error": str(e)}, status=400)
    

@method_decorator(csrf_exempt, name='dispatch')
class generate_cv(APIView):
    permission_classes = [IsAuthenticated]
    def post(self,request):
        fields = request.data
        user = request.user

        if not user.is_authenticated:
            return Response({"error": "User not authenticated"}, status=401)

        profile = get_object_or_404(Student, id__user=user)

        # Initialize PDF
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            topMargin=20,
            leftMargin=30,  # Reduced left margin
            rightMargin=30,  # Reduced right margin
        )

        # Styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name="Title",
            parent=styles["Title"],
            fontSize=18,
            leading=22,
            spaceAfter=10,
            alignment=1,  # Center alignment
        )
        section_header_style = ParagraphStyle(
            name="SectionHeader",
            parent=styles["Heading2"],
            fontSize=14,
            leading=18,
            spaceAfter=4,  # Reduce space after section header
            textColor=colors.HexColor("#c43119"),  # Custom color
        )
        body_style = styles["BodyText"]
        body_style.fontSize = 11
        body_style.leading = 14

        # Helper to format dates
        def format_date(date):
            return date.strftime("%d %B %Y") if date else "Ongoing"

        # Content container
        content = []

        # Add dynamic sections with optional bullet points
        def add_section(title, queryset, formatter, bullet_points=False):
            content.append(Paragraph(title, section_header_style))
            # Add a horizontal line under the section header
            line = Line(0, 0, 500, 0, strokeColor=colors.HexColor("#c43119"))  # Adjust line length
            drawing = Drawing(500, 1)  # Adjust drawing width
            drawing.add(line)
            content.append(drawing)
            content.append(Spacer(1, 4))  # Reduce space between header and content

            if bullet_points:
                # Add items as a bulleted list
                items = [Paragraph(formatter(obj), body_style) for obj in queryset]
                content.append(ListFlowable(
                    [ListItem(i) for i in items],
                    bulletType="bullet",  # Use bullets
                    start="circle",  # Specify bullet style
                    leftIndent=10,  # Indent for bullets
                    bulletFontSize=8,  # Decrease bullet size
                    bulletOffset=10,  # Increase space between bullet and text
                ))
            else:
                # Add items normally
                for obj in queryset:
                    content.append(Paragraph(formatter(obj), body_style))
            content.append(Spacer(1, 8))  # Reduce space between sections

        # Title
        content.append(Paragraph(f"{user.get_full_name()}", title_style))
        content.append(Spacer(1, 8))  # Reduce space after title

        # Dynamic Sections
        if fields.get("achievements", False):
            achievements = Achievement.objects.filter(unique_id=profile)
            add_section(
                "Achievements",
                achievements,
                lambda a: f"{a.achievement} ({a.achievement_type}) - {a.issuer} ({format_date(a.date_earned)})",
                bullet_points=True,
            )

        if fields.get("education", False):
            education = Education.objects.filter(unique_id=profile)
            add_section(
                "Education",
                education,
                lambda e: f"{e.degree} in {e.stream or 'General'} from {e.institute}, Grade: {e.grade} ({format_date(e.sdate)} - {format_date(e.edate)})"
            )

        if fields.get("skills", False):
            skills = Has.objects.filter(unique_id=profile)
            add_section(
                "Skills",
                skills,
                lambda s: f"{s.skill_id.skill} (Rating: {s.skill_rating}%)",
                bullet_points=True,
            )

        if fields.get("experience", False):
            experience = Experience.objects.filter(unique_id=profile)
            add_section(
                "Experience",
                experience,
                lambda e: f"<b>{e.title}</b> at {e.company} ({format_date(e.sdate)} - {format_date(e.edate)})<br/>{e.description or 'No description'}"
            )

        if fields.get("projects", False):
            projects = Project.objects.filter(unique_id=profile)
            add_section(
                "Projects",
                projects,
                lambda p: f"<b>{p.project_name}</b><br/>{p.summary or 'No description'} (Status: {p.project_status})",
                bullet_points=True,
            )

        if fields.get("courses", False):
            courses = Course.objects.filter(unique_id=profile)
            add_section(
                "Courses",
                courses,
                lambda c: f"{c.course_name} - {c.description or 'No description'} (License: {c.license_no or 'N/A'})",
                bullet_points=True,
            )

        # Build the PDF
        doc.build(content)
        buffer.seek(0)

        # Response
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response['Content-Disposition'] = 'attachment; filename="resume.pdf"'
        return response

@permission_classes([IsAuthenticated]) 
class ApplyForPlacement(APIView):
    def post(self,request):
        user = request.user
        profile = get_object_or_404(ExtraInfo, user=user)
        student = Student.objects.get(id_id=profile.id)
        placement_id = request.data.get('jobId')
        placement = PlacementSchedule.objects.get(id=placement_id)
        print(f"User: {user}, Profile: {profile}, Student: {student}, Placement ID: {placement_id}") 

        try:
            application = StudentApplication.objects.create(
                schedule_id = placement,
                unique_id = student,
                current_status = "accept",
            )


            return JsonResponse({"message": "Successfully Applied"}, status=201)

        except Exception as e:
            print(f"Error creating application: {str(e)}")
            return JsonResponse({"error": str(e)}, status=400)
        

    def get(self, request,id):
        schedule = get_object_or_404(PlacementSchedule, id=id)  
        applications = StudentApplication.objects.filter(schedule_id_id=schedule.id)
        print(schedule.id)

        students_data = []
        for application in applications:
            roll_no = application.unique_id_id
            print(roll_no)
            student = get_object_or_404(Student, id_id=roll_no)
            user = get_object_or_404(User, username=roll_no)

            students_data.append({
                'id':application.id,
                'name': f"{user.first_name} {user.last_name}",
                'roll_no': roll_no,
                'email': user.email,
                'cpi': student.cpi,
                'status': application.current_status,
            })
        print(students_data)
        return Response({'students': students_data}, status=200)
    
    def put(self, request, id):
        application = get_object_or_404(StudentApplication, id=id)

        new_status = request.data.get('status')
        if new_status is None:  
            return JsonResponse({"error": "Status is required"}, status=400)

        try:
            application.current_status = new_status
            application.save()
            print('changed')
            return JsonResponse({"message": "Status updated successfully"}, status=200)

        except Exception as e:
            print(f"Error updating application status: {str(e)}")
            return JsonResponse({"error": str(e)}, status=400)
    

@permission_classes([IsAuthenticated])
class NextRoundDetails(APIView):
    def post(self,request,id):
        round_no = request.data.get('round_no')
        test_type = request.data.get('test_type')
        test_date = request.data.get('test_date')
        description = request.data.get('description')

        try:
            next_round = NextRoundInfo.objects.create(
                schedule_id_id = id,
                round_no = round_no,
                test_type = test_type,
                test_date = test_date,
                description = description,
            )
            return JsonResponse({"message": "Successfully Created"}, status=201)

        except Exception as e:
            print(f"Error creating round: {str(e)}")
            return JsonResponse({"error": str(e)}, status=400)
        
    def get(self,request):
        user = request.user
        next_data=[]
        print(user.username)
        if user.username=='omvir' or user.username=='anilk':
            next_round_data = NextRoundInfo.objects.all()
            for nr in next_round_data:
                try:
                    schedule = PlacementSchedule.objects.get(id=nr.schedule_id_id)
                    print("Valid Schedule:", schedule.title) 
                except PlacementSchedule.DoesNotExist:
                    print("No schedule found for schedule_id:", nr.schedule_id_id)
                next_data.append({
                    'id':nr.schedule_id_id,
                    'company_name':schedule.title,
                    'date':nr.test_date,
                    'type':nr.test_type,
                    'round':nr.round_no,
                    'description':nr.description,
                })

        else:
            profile = get_object_or_404(ExtraInfo, user=user)
            roll_no = profile.id
            applications = StudentApplication.objects.filter(unique_id_id=roll_no,current_status='accept')
        
            for application in applications:
                next_round_data = NextRoundInfo.objects.filter(schedule_id=application.schedule_id)
                for nr in next_round_data:
                    next_data.append({
                        'id':nr.schedule_id_id,
                        'company_name':application.schedule_id.title,
                        'date':nr.test_date,
                        'type':nr.test_type,
                        'round':nr.round_no,
                        'description':nr.description,
                    })
        
        return Response({'schedule_data': next_data}, status=200)
    

    def put(self, request, round_id):
        next_round = get_object_or_404(NextRoundInfo, id=round_id)

        round_no = request.data.get('round_no')
        test_type = request.data.get('test_type')
        test_date = request.data.get('test_date')
        description = request.data.get('description')

        try:
            if round_no is not None:
                next_round.round_no = round_no
            if test_type is not None:
                next_round.test_type = test_type
            if test_date is not None:
                next_round.test_date = test_date
            if description is not None:
                next_round.description = description

            next_round.save()

            return JsonResponse({"message": "Successfully Updated"}, status=200)

        except Exception as e:
            print(f"Error updating round: {str(e)}")
            return JsonResponse({"error": str(e)}, status=400)



@permission_classes([IsAuthenticated])
class TrackStatus(APIView):
    def get(self,request,id):
        user = request.user
        profile = get_object_or_404(ExtraInfo, user=user)
        roll_no = profile.id
        status='reject'
        if user.username!='omvir' and user.username!='anilk':
            application = StudentApplication.objects.get(unique_id_id=roll_no, schedule_id_id=id)
            status = application.current_status
        data = []


        if user.username=='omvir' or user.username=='anilk' or status != 'reject':
            rounds = NextRoundInfo.objects.filter(schedule_id_id=id).order_by('round_no')
            round_count = rounds.count()

            if round_count==0:
                data.append({
                    'round_no': 0,
                    'test_name': 'Yet to be updated',
                })
            
            else:
                for round_info in rounds[:round_count - 1]: 
                    data.append({
                        'round_no': round_info.round_no,
                        'test_name': round_info.test_type,
                    })

                if round_count > 0:
                    last_round_info = rounds[round_count - 1]
                    data.append({
                        'round_no': last_round_info.round_no,
                        'test_name': last_round_info.test_type,
                        'test_date': last_round_info.test_date,
                        'description': last_round_info.description,
                    })
        
        else:
            data.append({
                'round_no':-1,
            })

        return Response({'next_data': data}, status=200)

@permission_classes([IsAuthenticated])
class DownloadApplications(APIView):
    def get(self, request, id):
        schedule = get_object_or_404(PlacementSchedule, id=id)
        applications = StudentApplication.objects.filter(schedule_id_id=schedule.id)

        wb = Workbook()
        ws = wb.active
        ws.title = "Applications"

        headers = ['ID', 'Name', 'Roll Number', 'Email', 'CPI', 'Status']
        ws.append(headers)

        for application in applications:
            roll_no = application.unique_id_id
            student = get_object_or_404(Student, id_id=roll_no)
            user = get_object_or_404(User, username=roll_no)

            row = [
                application.id,
                f"{user.first_name} {user.last_name}",
                roll_no,
                user.email,
                student.cpi,
                application.current_status,
            ]
            ws.append(row)

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="applications_{schedule.title}.xlsx"'

        wb.save(response)
        return response


@permission_classes([IsAuthenticated])
class DownloadStatistics(APIView):
    def get(self, request):
        student_records = StudentRecord.objects.all()

        if not student_records.exists():
            return Response({"error": "No student records found"}, status=status.HTTP_404_NOT_FOUND)

        wb = Workbook()
        ws = wb.active
        ws.title = "Placement Statistics"

        headers = ['First Name', 'Placement Name', 'Batch', 'Branch', 'CTC', 'Year']
        ws.append(headers)

        for student in student_records:
            try:
                cur_student = Student.objects.get(id_id=student.unique_id_id)
                cur_placement = PlacementRecord.objects.get(id=student.record_id_id)
                user = User.objects.get(username=student.unique_id_id)

                row = [
                    user.first_name,
                    cur_placement.name,
                    cur_placement.year,
                    cur_student.specialization,
                    cur_placement.ctc,
                    cur_placement.year,
                ]
                ws.append(row)

            except (Student.DoesNotExist, PlacementRecord.DoesNotExist, User.DoesNotExist) as e:
                continue

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="placement_statistics.xlsx"'

        wb.save(response)
        return response


@permission_classes([IsAuthenticated])
class DebarStudents(APIView):
    def get(self,request):
        debared_students = DebarStudentInfo.objects.all()
        data = []

        for stud in debared_students:
            user = User.objects.get(username=stud.unique_id_id)
            data.append({
                'roll_no':stud.unique_id_id,
                'name':user.first_name,
                'description':stud.description
            })

        return Response(data, status=status.HTTP_200_OK)
    


@permission_classes([IsAuthenticated])
class DebaredDetails(APIView):
    def get(self,request,id):
        try:
            print("entered")
            user = User.objects.get(username=id)
            extra_info = ExtraInfo.objects.get(id=id)
            department = DepartmentInfo.objects.get(id=extra_info.department_id)
            student = Student.objects.get(id_id=id)
            debar = DebarStudentInfo.objects.filter(unique_id_id=id).count()
            print("entered")
            

            data = {
                'roll_no':user.username,
                'name':user.first_name,
                'email':user.email,
                'department':department.name,
                'year': datetime.now().year - student.batch,
                'programme':student.programme,
                'debar_status':debar
            }
            print(data)
            
            return Response(data,status=status.HTTP_200_OK)
        

        except User.DoesNotExist:
            return Response({"error": f"Student with id {id} not found"}, status=status.HTTP_404_NOT_FOUND)
        
    
    def post(self,request,id):
        debar = DebarStudentInfo.objects.filter(unique_id_id=id).count()
        # print(debar,id)
        if debar==1:
            return Response("already present",status=status.HTTP_300_MULTIPLE_CHOICES)
        
        roll_no=id
        description=request.data.get('reason')

        try:
            obj = DebarStudentInfo.objects.create(
                unique_id_id=id,
                description=description,
            )

            return Response("Successfully debared",status=status.HTTP_200_OK)
        
        except Exception as e:
            print(e)
            return Response("Failed to debar",status=status.HTTP_300_MULTIPLE_CHOICES)
        
    def delete(self,request,id):
        try:
            debar = DebarStudentInfo.objects.get(unique_id_id=id)
            debar.delete()

            return Response("succesfully deleted",status=status.HTTP_200_OK)
        
        except Exception as e:
            return Response("error occured : "+ e )